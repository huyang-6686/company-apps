from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "historical documents"

def build(src_name, out_name):
    source = (OUT / src_name).read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        styles[style_name].font.name = "Microsoft YaHei"
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.bold = True
        styles[style_name].font.color.rgb = __import__("docx").shared.RGBColor(0, 0, 0)

    first_title = True
    in_code = False
    for raw in source.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if not line:
            doc.add_paragraph()
            continue
        if first_title and line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(line[2:].strip())
            first_title = False
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            continue
        if line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.add_run(line[2:].strip()).italic = True
            continue
        if line.startswith("|"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.15)
            p.add_run(line).font.name = "Consolas"
            p.runs[0].font.size = Pt(8.5)
            continue
        p = doc.add_paragraph()
        if in_code:
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        else:
            p.add_run(line)
    doc.save(OUT / out_name)

build("组织架构管理规范-20261019正式版.md", "组织架构管理规范-20261019正式版.docx")
build("应用搭建结构规范-20261019更新版正式版.md", "应用搭建结构规范-20261019更新版正式版.docx")
build("网盘映射管理与文件命名规范-20261019更新版正式版.md", "网盘映射管理与文件命名规范-20261019更新版正式版.docx")
build("任务发布与管理规范-20261020正式版.md", "任务发布与管理规范-20261020正式版.docx")
build("任务接口合同-20261020正式版.md", "任务接口合同-20261020正式版.docx")
build("统一信息通知与应用消息通道规范-20261020正式版.md", "统一信息通知与应用消息通道规范-20261020正式版.docx")
build("五份规范一致性自检记录-20261020.md", "五份规范一致性自检记录-20261020.docx")
build("各应用最终规范发布与接入执行提示词-20261020.md", "各应用最终规范发布与接入执行提示词-20261020.docx")
